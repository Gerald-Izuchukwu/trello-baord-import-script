#!/usr/bin/env python3
"""
Create the Servana DevOps Trello board straight from the Excel plan.

Creates: the board, every list, every label (with colours) and every card
(description, labels, due date). No manual list or label setup needed.

Setup (once):
    pip install requests openpyxl

Get Trello credentials:
    1. API key:  https://trello.com/power-ups/admin  (create a Power-Up, open
       "API key", copy the key). Trello's admin pages change, so if the link
       moves, search "Trello API key" in Trello help.
    2. Token:    open this URL in a browser while logged in, click Allow, and
       copy the token it shows (replace YOUR_KEY):
       https://trello.com/1/authorize?expiration=1day&scope=read,write&response_type=token&name=ServanaBoard&key=YOUR_KEY

Run:
    export TRELLO_KEY=your_key            (Windows PowerShell: $env:TRELLO_KEY="your_key")
    export TRELLO_TOKEN=your_token        (Windows PowerShell: $env:TRELLO_TOKEN="your_token")
    python create_trello_board.py --dry-run     # checks the file, makes no changes
    python create_trello_board.py               # creates the board

Keep the key and token private. Never commit them to Git.
"""
import argparse
import getpass
import os
import sys
import time
from datetime import timedelta, timezone

import requests
from openpyxl import load_workbook

API = "https://api.trello.com/1"
SHEET = "Trello Import"

# Label name -> colour candidates, tried in order. None = no colour (grey in Trello).
# If Trello rejects a colour, the next candidate is tried.
LABELS = [
    ("Gerald", ["blue"]),
    ("Victor", ["purple"]),
    ("Urgent", ["orange"]),
    ("Critical", ["red"]),
    ("Overdue", ["red_dark", "red"]),
    ("Not Started", [None]),
    ("Started", ["yellow"]),
    ("Incomplete", ["pink"]),
    ("Completed", ["green"]),
    ("Blocked", ["black"]),
    ("Milestone", ["lime"]),
    ("Check-in", ["sky"]),
    ("Stretch", [None]),
    ("Guide", ["sky_light", "sky"]),
    ("Infra", [None]),
    ("Docker", [None]),
    ("CI/CD", [None]),
    ("Security", [None]),
    ("Monitoring", [None]),
    ("Integrations", [None]),
    ("Testing", [None]),
    ("Docs", [None]),
    ("Demo", [None]),
]


class TrelloError(Exception):
    pass


class Trello:
    def __init__(self, key, token):
        self.auth = {"key": key, "token": token}
        self.session = requests.Session()

    def call(self, method, path, check=True, **fields):
        """POST/PUT fields go in the body (long descriptions are safe)."""
        for attempt in range(6):
            time.sleep(0.15)  # stay well under Trello's rate limit
            if method == "GET":
                r = self.session.get(API + path, params={**self.auth, **fields}, timeout=30)
            else:
                r = self.session.request(method, API + path, params=self.auth, data=fields, timeout=30)
            if r.status_code == 429:
                time.sleep(10)
                continue
            if r.status_code >= 500:
                time.sleep(2 * (attempt + 1))
                continue
            break
        if check and not r.ok:
            raise TrelloError(f"{method} {path} failed ({r.status_code}): {r.text[:300]}")
        return r


def read_rows(path, tz_hours):
    wb = load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"Sheet '{SHEET}' not found in {path}")
    ws = wb[SHEET]
    tz = timezone(timedelta(hours=tz_hours))
    rows = []
    for r in ws.iter_rows(min_row=2, max_col=5, values_only=True):
        name, desc, lst, labels, due = r
        if not name:
            continue
        due_iso = None
        if due:
            due_iso = due.replace(tzinfo=tz).astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")
        rows.append({
            "name": str(name).strip(),
            "desc": str(desc or ""),
            "list": str(lst).strip(),
            "labels": [x.strip() for x in str(labels or "").split(",") if x.strip()],
            "due": due_iso,
        })
    return rows


def main():
    ap = argparse.ArgumentParser(description="Create the Servana DevOps Trello board from the Excel plan.")
    ap.add_argument("--file", default="Servana-DevOps-Trello-Board.xlsx", help="path to the Excel plan")
    ap.add_argument("--name", default="Servana DevOps Board", help="board name")
    ap.add_argument("--workspace", default=None, help="Trello workspace (organization) ID; omit for your personal boards")
    ap.add_argument("--tz-offset", type=float, default=1, help="hours from UTC for the due times in the sheet (Lagos = 1)")
    ap.add_argument("--dry-run", action="store_true", help="read the file and print what would be created")
    args = ap.parse_args()

    rows = read_rows(args.file, args.tz_offset)

    lists = []
    for r in rows:
        if r["list"] not in lists:
            lists.append(r["list"])
    known = {n for n, _ in LABELS}
    unknown = sorted({lb for r in rows for lb in r["labels"]} - known)
    if unknown:
        sys.exit(f"Labels in the sheet with no colour defined in this script: {unknown}. Add them to LABELS.")

    print(f"Plan: {len(rows)} cards, {len(lists)} lists, {len(LABELS)} labels.")
    if args.dry_run:
        for l in lists:
            print(f"  {sum(1 for r in rows if r['list'] == l):3d} cards  {l}")
        print("Dry run only. Nothing was created.")
        return

    key = os.environ.get("TRELLO_KEY") or getpass.getpass("Trello API key: ")
    token = os.environ.get("TRELLO_TOKEN") or getpass.getpass("Trello token: ")
    t = Trello(key.strip(), token.strip())

    # Confirm the credentials work before creating anything.
    me = t.call("GET", "/members/me", fields="username").json()
    print(f"Signed in as {me.get('username')}")

    board_fields = {"name": args.name, "defaultLists": "false", "defaultLabels": "false",
                    "prefs_permissionLevel": "private"}
    if args.workspace:
        board_fields["idOrganization"] = args.workspace
    board = t.call("POST", "/boards", **board_fields).json()
    print(f"Board created: {board['url']}")
    try:
        list_ids = {}
        for l in lists:
            list_ids[l] = t.call("POST", "/lists", name=l, idBoard=board["id"], pos="bottom").json()["id"]
        print(f"Created {len(list_ids)} lists")

        label_ids = {}
        for name, candidates in LABELS:
            for color in candidates:
                fields = {"name": name, "idBoard": board["id"]}
                if color:
                    fields["color"] = color
                r = t.call("POST", "/labels", check=False, **fields)
                if r.ok:
                    label_ids[name] = r.json()["id"]
                    break
            else:
                raise TrelloError(f"Could not create label '{name}'")
        print(f"Created {len(label_ids)} labels")

        for i, r in enumerate(rows, 1):
            fields = {
                "idList": list_ids[r["list"]],
                "name": r["name"],
                "desc": r["desc"],
                "pos": "bottom",
                "idLabels": ",".join(label_ids[x] for x in r["labels"]),
            }
            if r["due"]:
                fields["due"] = r["due"]
            t.call("POST", "/cards", **fields)
            if i % 10 == 0 or i == len(rows):
                print(f"  cards {i}/{len(rows)}")
    except Exception as e:
        print(f"\nStopped part-way: {e}")
        print(f"Delete the half-built board and run again: {board['url']}")
        sys.exit(1)

    print(f"\nDone. Open your board: {board['url']}")


if __name__ == "__main__":
    main()